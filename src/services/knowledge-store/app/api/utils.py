import os
import shutil

import pdfplumber
import pytesseract
from docx import Document
from fastapi import UploadFile
from pdf2image import convert_from_path


async def docx_to_text(filepath: str) -> str:
    """
    Extract text content from a DOCX file.

    Iterates through paragraphs and tables to compile the full text.

    Args:
        filepath (str): The absolute or relative path to the .docx file.

    Returns:
        str: The extracted text content.
    """
    doc = Document(filepath)
    text = []
    for paragraph in doc.paragraphs:
        text.append(paragraph.text)

    for table in doc.tables:
        for row in table.rows:
            row_text = [cell.text for cell in row.cells]
            text.append('\t'.join(row_text))

    return '\n'.join(text)

async def image_pdf_to_text(filepath: str) -> str:
    """
    Convert a scanned/image-based PDF file to text using OCR (Tesseract).

    Args:
        filepath (str): The path to the PDF file.

    Returns:
        str: The extracted text from OCR.
    """
    images = convert_from_path(filepath)
    text = '\n'.join([pytesseract.image_to_string(img) for img in images])

    return text

async def pdf_to_text(filepath: str) -> str:
    """
    Extract text from a text-based PDF file.

    Args:
        filepath (str): The path to the PDF file.

    Returns:
        str: The extracted text.
    """
    text = ''
    with pdfplumber.open(filepath) as pdf:
        for page in pdf.pages:
            text += page.extract_text() + '\n'
    return text

async def pdf_pages(filepath: str) -> list[str]:
    with pdfplumber.open(filepath) as pdf:
        return [page.extract_text() or '' for page in pdf.pages]

async def image_pdf_pages(filepath: str) -> list[str]:
    return [pytesseract.image_to_string(img) for img in convert_from_path(filepath)]

async def save_file(file: UploadFile, dir_name: str) -> str:
    """
    Save an uploaded file to the local filesystem.

    Args:
        file (UploadFile): The file object uploaded via FastAPI.
        dir_name (str): The subdirectory name within 'documents/' to save the file.

    Returns:
        str: The full path to the saved file.
    """
    path = os.path.join('documents/', dir_name, file.filename)
    os.makedirs(os.path.dirname(path), exist_ok=True)
    with open(path, 'wb') as buffer:
        shutil.copyfileobj(file.file, buffer)
    return path
